"""Xuất báo cáo: HTML, XLSX, CSV, JSON."""
import csv
import html
import json
from collections import Counter, OrderedDict
from pathlib import Path

COLORS = {"PASS": "1B8A3A", "FAIL": "C0392B", "ERROR": "D68910", "SKIPPED": "7F8C8D"}
FIELDS = ["id", "module", "title", "type", "priority", "scenario", "input1", "input2", "expected",
          "status", "actual", "url", "error", "duration", "attempts", "screenshot"]


def summarize(results):
    c = Counter(r["status"] for r in results)
    total = len(results)
    executed = total - c.get("SKIPPED", 0)
    by_module = OrderedDict()
    for r in results:
        m = by_module.setdefault(r["module"] or "-", Counter())
        m[r["status"]] += 1
        m["total"] += 1
    return {
        "total": total, "executed": executed,
        "pass": c.get("PASS", 0), "fail": c.get("FAIL", 0),
        "error": c.get("ERROR", 0), "skipped": c.get("SKIPPED", 0),
        "pass_rate": round(100 * c.get("PASS", 0) / executed, 1) if executed else 0.0,
        "duration": round(sum(r["duration"] for r in results), 1),
        "by_module": {k: dict(v) for k, v in by_module.items()},
    }


def write_all(run_dir, meta, results):
    run_dir = Path(run_dir)
    summ = summarize(results)
    (run_dir / "results.json").write_text(
        json.dumps({"meta": meta, "summary": summ, "results": results}, ensure_ascii=False, indent=2),
        encoding="utf-8")
    with open(run_dir / "results.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, FIELDS, extrasaction="ignore")
        w.writeheader()
        w.writerows(results)
    _xlsx(run_dir / "results.xlsx", meta, summ, results)
    _html(run_dir / "report.html", meta, summ, results)
    return summ


def _xlsx(path, meta, summ, results):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Tổng quan"
    rows = [("Báo cáo kiểm thử tự động Khan Academy", ""), ("Thời gian chạy", meta.get("started", "")),
            ("Website", meta.get("base_url", "")), ("Trình duyệt", meta.get("browser", "")),
            ("Bộ test case", meta.get("file", "")), ("", ""),
            ("Tổng test case", summ["total"]), ("Đã thực thi", summ["executed"]),
            ("PASS", summ["pass"]), ("FAIL", summ["fail"]), ("ERROR", summ["error"]),
            ("SKIPPED", summ["skipped"]), ("Tỷ lệ PASS (%)", summ["pass_rate"]),
            ("Tổng thời gian (s)", summ["duration"])]
    for r in rows:
        ws.append(r)
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(["Module", "Tổng", "PASS", "FAIL", "ERROR", "SKIPPED"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for m, v in summ["by_module"].items():
        ws.append([m, v.get("total", 0), v.get("PASS", 0), v.get("FAIL", 0), v.get("ERROR", 0), v.get("SKIPPED", 0)])
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 40

    ws2 = wb.create_sheet("Kết quả chi tiết")
    ws2.append(FIELDS)
    for c in ws2[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2C3E50")
    for r in results:
        ws2.append([r.get(k, "") for k in FIELDS])
        cell = ws2.cell(row=ws2.max_row, column=FIELDS.index("status") + 1)
        cell.fill = PatternFill("solid", fgColor=COLORS.get(r["status"], "FFFFFF"))
        cell.font = Font(bold=True, color="FFFFFF")
    widths = dict(id=9, module=14, title=44, type=12, priority=9, scenario=16, input1=26, input2=14, expected=34,
                  status=10, actual=60, url=40, error=40, duration=9, attempts=9, screenshot=26)
    for i, k in enumerate(FIELDS, 1):
        ws2.column_dimensions[ws2.cell(row=1, column=i).column_letter].width = widths.get(k, 15)
    for row in ws2.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = ws2.dimensions
    wb.save(path)


def _html(path, meta, s, results):
    """Báo cáo HTML độc lập: dùng chung theme.css với app, có nút Sáng/Tối."""
    e = html.escape
    css = (Path(__file__).resolve().parent.parent / "web" / "static" / "theme.css").read_text(encoding="utf-8")
    icon = {"PASS": "✓", "FAIL": "✕", "ERROR": "!", "SKIPPED": "–"}

    def dist(c):
        return '<div class="dist">' + "".join(
            f'<i class="{k}" style="flex:{c.get(k, 0)}" title="{k}: {c.get(k, 0)}"></i>'
            for k in ("PASS", "FAIL", "ERROR", "SKIPPED") if c.get(k, 0)) + "</div>"

    rows = []
    for r in results:
        shot = (f'<a href="{e(r["screenshot"])}" target="_blank"><img class="shot" src="{e(r["screenshot"])}" '
                f'alt="Ảnh {e(r["id"])}"></a>') if r["screenshot"] else ""
        rows.append(
            f'<tr data-s="{r["status"]}"><td class="mono">{e(r["id"])}</td>'
            f'<td>{e(r["title"])}<div class="muted small">{e(r["module"])} · <span class="tag">{e(r.get("type",""))}</span> '
            f'<span class="pri {e(r["priority"])}">{e(r["priority"])}</span></div></td>'
            f'<td class="mono small" style="max-width:260px">{e(r["input1"][:120])}{'…' if len(r['input1'])>120 else ''}{(" / " + e(r["input2"])) if r["input2"] else ""}</td>'
            f'<td class="mono small">{e(r["expected"])}</td>'
            f'<td><span class="pill {r["status"]}">{icon[r["status"]]} {r["status"]}</span></td>'
            f'<td class="small">{e(r["error"])}<div class="muted">{e((r["actual"] or "")[:200])}</div></td>'
            f'<td class="num">{r["duration"]}s</td><td>{shot}</td></tr>')
    mods = "".join(
        f'<div class="mod"><span>{e(m)}</span>{dist(v)}<span class="num muted">{v.get("PASS",0)}/{v.get("total",0)-v.get("SKIPPED",0)}</span></div>'
        for m, v in s["by_module"].items())
    cnt = {k: s[k.lower()] for k in ("PASS", "FAIL", "ERROR", "SKIPPED")}
    legend = "".join(f'<span><i style="background:var(--{"skip" if k=="SKIPPED" else k.lower()})"></i>{k} {v}</span>'
                     for k, v in cnt.items())
    chips = "".join(f'<button class="chip" data-f="{k}" aria-pressed="false">{k}</button>' for k in ("PASS", "FAIL", "ERROR", "SKIPPED"))
    path.write_text(f"""<!doctype html><html lang="vi"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>Báo cáo kiểm thử – Khan Academy ({e(meta.get('started',''))})</title>
<link href="https://fonts.googleapis.com/css2?family=Be+Vietnam+Pro:wght@400;500;600;700&family=JetBrains+Mono:wght@400;600&display=swap" rel="stylesheet">
<style>{css}
.wrap{{max-width:1240px;margin:0 auto;padding:20px 20px 48px;display:grid;grid-template-columns:minmax(0,1fr);gap:18px}}
.small{{font-size:12.5px}}.chips{{display:flex;gap:6px;flex-wrap:wrap}}
.chip{{border:1px solid var(--line);background:var(--surface);color:var(--muted);border-radius:99px;padding:3px 11px;font-size:12.5px}}
.chip[aria-pressed=true]{{background:var(--accent-soft);border-color:var(--accent);color:var(--accent);font-weight:600}}
h1{{font-size:22px;letter-spacing:-.01em}}.meta{{display:flex;gap:6px 18px;flex-wrap:wrap;color:var(--muted);font-size:13px;margin-top:6px}}
</style></head><body>
<header class="top"><div class="top-in"><div class="brand"><span class="mark"><svg width="16" height="16" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="3" stroke-linecap="round" stroke-linejoin="round"><path d="M4 12.5l5 5L20 6.5"/></svg></span>Khan Academy AutoTest</div>
<span class="spacer"></span><div class="seg" role="group" aria-label="Giao diện"><button data-th="system">Tự động</button><button data-th="light">Sáng</button><button data-th="dark">Tối</button></div></div></header>
<main class="wrap">
<div><h1>Báo cáo kiểm thử tự động</h1>
<div class="meta"><span>Chạy lúc <b class="num">{e(meta.get('started',''))}</b></span><span>Website <b>{e(meta.get('base_url',''))}</b></span>
<span>{e(meta.get('browser',''))}</span><span>Bộ test: {e(meta.get('file',''))}</span><span>Tổng thời gian <b class="num">{s['duration']}s</b></span></div></div>
<section class="card"><div class="card-b">
<div class="kpis"><div class="kpi"><b>{s['total']}</b><span>Tổng</span></div><div class="kpi PASS"><b>{s['pass']}</b><span>PASS</span></div>
<div class="kpi FAIL"><b>{s['fail']}</b><span>FAIL</span></div><div class="kpi ERROR"><b>{s['error']}</b><span>ERROR</span></div>
<div class="kpi SKIPPED"><b>{s['skipped']}</b><span>SKIPPED</span></div><div class="kpi"><b>{s['pass_rate']}%</b><span>Tỷ lệ PASS ({s['executed']} đã chạy)</span></div></div>
<div style="margin-top:16px">{dist(cnt)}<div class="legend">{legend}</div></div>
<h3 style="font-size:13px;margin:18px 0 6px">Theo module</h3>{mods}</div></section>
<section class="card"><div class="card-h"><h2>Chi tiết từng test case</h2><span class="spacer"></span>
<div class="chips"><button class="chip" data-f="" aria-pressed="true">Tất cả</button>{chips}</div></div>
<div class="tbl-wrap"><table><thead><tr><th>ID</th><th>Mô tả</th><th>Dữ liệu vào</th><th>Kỳ vọng</th><th>Kết quả</th><th>Thực tế / lỗi</th><th>Thời gian</th><th>Ảnh</th></tr></thead>
<tbody id="rows">{''.join(rows)}</tbody></table></div></section>
</main>
<script>
function setTheme(t){{try{{localStorage.setItem('khan-theme',t)}}catch(e){{}}
 if(t==='system')document.documentElement.removeAttribute('data-theme');else document.documentElement.dataset.theme=t;
 document.querySelectorAll('.seg button').forEach(b=>b.setAttribute('aria-pressed',b.dataset.th===t))}}
document.querySelectorAll('.seg button').forEach(b=>b.onclick=()=>setTheme(b.dataset.th));
var s='system';try{{s=localStorage.getItem('khan-theme')||'system'}}catch(e){{}}setTheme(s);
document.querySelectorAll('.chip').forEach(c=>c.onclick=()=>{{
 document.querySelectorAll('.chip').forEach(x=>x.setAttribute('aria-pressed',x===c));
 document.querySelectorAll('#rows tr').forEach(r=>r.hidden=c.dataset.f&&r.dataset.s!==c.dataset.f)}});
</script></body></html>""", encoding="utf-8")
