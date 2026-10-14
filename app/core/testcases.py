"""Đọc test case (đầu vào) từ CSV / XLSX / JSON."""
import csv
import json
from pathlib import Path

COLUMNS = ["id", "module", "title", "type", "priority", "scenario", "input1", "input2",
           "expected_mode", "expected", "note"]
SCENARIOS = {
    "login": "Đăng nhập: input1=email, input2=mật khẩu",
    "login_then_open": "Đăng nhập bằng {EMAIL}/{PASSWORD} rồi mở trang input1 (cần tài khoản)",
    "login_enter": "Đăng nhập nhưng gửi form bằng phím Enter: input1=email, input2=mật khẩu",
    "click": "Mở trang input1, bấm phần tử CSS selector input2, rồi kiểm tra",
    "search": "Tìm kiếm: input1=từ khóa (gõ vào ô tìm kiếm + Enter)",
    "open_page": "Mở trang: input1=đường dẫn (vd /math)",
}
MODES = {
    "url_contains": "URL hiện tại chứa chuỗi (dùng | để OR)",
    "url_not_contains": "URL hiện tại KHÔNG chứa chuỗi",
    "text_contains": "Nội dung trang chứa chuỗi (không phân biệt hoa/thường, | để OR)",
    "text_not_contains": "Nội dung trang KHÔNG chứa chuỗi",
    "title_contains": "Tiêu đề tab chứa chuỗi",
    "element_visible": "Phần tử CSS selector hiển thị",
    "results_min": "Số kết quả tìm kiếm >= N",
}


def load(path):
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".csv":
        with open(path, newline="", encoding="utf-8-sig") as f:
            rows = list(csv.DictReader(f))
    elif ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        ws = load_workbook(path, read_only=True, data_only=True).active
        it = ws.iter_rows(values_only=True)
        head = [str(h or "").strip().lower() for h in next(it)]
        rows = [{h: ("" if v is None else str(v)) for h, v in zip(head, r)} for r in it]
    elif ext == ".json":
        rows = json.loads(path.read_text(encoding="utf-8"))
    else:
        raise ValueError(f"Định dạng không hỗ trợ: {ext} (dùng .csv, .xlsx, .json)")
    out = []
    for i, r in enumerate(rows, 1):
        r = {k.strip().lower(): ("" if v is None else str(v)) for k, v in r.items() if k}
        if not any(r.values()):
            continue
        tc = {c: r.get(c, "").strip() for c in COLUMNS}
        # input giữ nguyên khoảng trắng (test dấu cách trong email/mật khẩu)
        tc["input1"], tc["input2"] = r.get("input1", ""), r.get("input2", "")
        tc["id"] = tc["id"] or f"TC{i:03d}"
        tc["scenario"] = tc["scenario"].lower()
        tc["expected_mode"] = tc["expected_mode"].lower()
        tc["enabled"] = r.get("enabled", "1").strip().lower() not in ("0", "no", "false", "n")
        out.append(tc)
    return out


def validate(tcs):
    errs, seen = [], set()
    for t in tcs:
        if t["id"] in seen:
            errs.append(f"{t['id']}: trùng ID")
        seen.add(t["id"])
        if t["scenario"] not in SCENARIOS:
            errs.append(f"{t['id']}: scenario '{t['scenario']}' không hợp lệ")
        if t["expected_mode"] not in MODES:
            errs.append(f"{t['id']}: expected_mode '{t['expected_mode']}' không hợp lệ")
    return errs
